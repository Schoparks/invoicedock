"""Core end-to-end checks against a disposable local invoice library."""
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db, main


class InvoiceApiTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.data = Path(self.temporary.name) / 'data'
        for module in (db, main):
            patcher = patch.object(module, 'DATA', self.data)
            patcher.start()
            self.addCleanup(patcher.stop)
        environment = patch.dict(os.environ, {'INVOICE_MAIL_WORKER':'0','ALLOW_REGISTRATION':'true','COOKIE_SECURE':'false'})
        environment.start()
        self.addCleanup(environment.stop)
        self.client = TestClient(main.app, headers={'X-Requested-With': 'InvoiceDesk'})
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.user = self.call('POST', '/api/auth/register', {'email':'test@example.test','name':'测试用户','password':'test-password-123'})['user']

    def call(self, method, path, payload=None, status=200):
        response = self.client.request(method, path, json=payload)
        self.assertEqual(response.status_code, status, response.text)
        return response.json()

    def invoice(self, amount=1000, **fields):
        return self.call('POST', '/api/invoices', {
            'seller': '示例商家', 'buyer': '测试公司', 'date': '2026-10-01',
            'category': '餐饮', 'amount_cents': amount, 'verified': True, **fields,
        })

    def batch(self, invoices, **fields):
        return self.call('POST', '/api/batches', {
            'name': '十月报销', 'invoice_ids': [item['id'] for item in invoices], **fields,
        })

    def state(self):
        return self.call('GET', '/api/state')

    def match(self, target=1000, **fields):
        return self.call('POST', '/api/match', {'target_cents': target, **fields})

    def xml_import(self, number='001234567890', amount='12.34'):
        blob = (f'<Invoice><InvoiceNumber>{number}</InvoiceNumber>'
                '<SellerName>原始商家</SellerName><BuyerName>测试公司</BuyerName>'
                '<IssueDate>2026-10-01</IssueDate>'
                f'<TotalTaxIncludedAmount>{amount}</TotalTaxIncludedAmount></Invoice>').encode()
        response = self.client.post('/api/import', files=[('files', ('原始发票.xml', blob, 'application/xml'))])
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual(result['errors'], [])
        self.assertEqual(len(result['imported']), 1)
        return result['imported'][0], blob

    def test_import_preserves_original_deduplicates_and_requires_review(self):
        invoice, original = self.xml_import()
        self.assertEqual(invoice['number'], '001234567890')
        self.assertEqual(invoice['amount_cents'], 1234)
        self.assertFalse(invoice['verified'])
        self.assertEqual(self.client.get(invoice['file_url']).content, original)
        duplicate = self.client.post('/api/import', files=[('files', ('改名.xml', original, 'application/xml'))]).json()
        self.assertEqual(duplicate['duplicates'], ['改名.xml'])
        self.assertEqual(duplicate['imported'], [])
        self.assertEqual(len(self.state()['invoices']), 1)
        self.assertEqual(self.state()['summary']['pending_review'], 1)
        self.assertEqual(self.match(1234)['candidate_count'], 0)
        self.call('POST', '/api/batches', {'name': '不可报销', 'invoice_ids': [invoice['id']]}, status=409)
        self.call('PATCH', '/api/invoices/' + invoice['id'], {'verified': True})
        match = self.match(1234)['combinations'][0]
        self.assertEqual(match['invoice_ids'], [invoice['id']])
        self.assertEqual((match['total_cents'], match['gap_cents']), (1234, 0))

    def test_exact_match_and_draft_reservation_prevent_double_use(self):
        self.invoice(700)
        left, right = self.invoice(510), self.invoice(490)
        result = self.match()['combinations'][0]
        self.assertEqual(set(result['invoice_ids']), {left['id'], right['id']})
        self.assertEqual(result['total_cents'], 1000)
        self.batch([left, right])
        self.call('POST', '/api/batches', {'name': '重复报销', 'invoice_ids': [left['id']]}, status=409)
        self.call('PATCH', '/api/invoices/' + left['id'], {'amount_cents': 520}, status=409)
        remaining = self.match()
        self.assertEqual(remaining['candidate_count'], 1)
        self.assertEqual(remaining['combinations'], [])
        self.assertIn('额度不足', remaining['message'])

    def test_match_closest_above_target_and_respects_quota(self):
        self.invoice(700)
        self.invoice(400)
        self.invoice(1500)
        matching = self.match()
        self.assertEqual(matching['combinations'][0]['total_cents'], 1100)
        self.assertEqual(matching['combinations'][0]['gap_cents'], -100)
        self.assertTrue(all(row['total_cents'] >= 1000 for row in matching['combinations']))
        quota = self.call('POST', '/api/quotas', {'month': '2026-10', 'category': '餐饮', 'amount_cents': 1200})
        matching = self.match(quota_id=quota['id'])
        self.assertEqual([row['total_cents'] for row in matching['combinations']], [1100])
        low_quota = self.call('POST', '/api/quotas', {'month': '2026-11', 'category': '餐饮', 'amount_cents': 900})
        error = self.call('POST', '/api/match', {'target_cents': 1000, 'quota_id': low_quota['id']}, status=400)
        self.assertIn('额度不足', error['detail'])
        combination = matching['combinations'][0]
        batch = self.call('POST', '/api/batches', {'name': '超出目标但在额度内',
                         'invoice_ids': combination['invoice_ids'], 'quota_id': quota['id']})
        self.assertEqual(batch['total_cents'], 1100)

    def test_match_below_prefers_exact_and_respects_remaining_quota(self):
        self.invoice(700)
        self.invoice(299)
        exact = self.invoice(1000)
        quota_invoice = self.invoice(600)
        matching = self.match(direction='below')
        self.assertEqual(matching['direction'], 'below')
        self.assertEqual(matching['combinations'][0]['invoice_ids'], [exact['id']])
        self.assertEqual(matching['combinations'][0]['gap_cents'], 0)
        self.assertTrue(all(row['total_cents'] <= 1000 for row in matching['combinations']))
        self.assertEqual(self.match()['direction'], 'above')
        self.batch([exact])
        self.assertEqual(self.match(direction='below')['combinations'][0]['total_cents'], 999)
        quota = self.call('POST', '/api/quotas', {'month': '2026-10', 'category': '餐饮', 'amount_cents': 600})
        matching = self.match(direction='below', quota_id=quota['id'])
        self.assertEqual(matching['target_cents'], 1000)
        self.assertEqual(matching['combinations'][0]['total_cents'], 600)
        self.assertEqual(matching['combinations'][0]['gap_cents'], 400)
        self.assertTrue(all(row['total_cents'] <= 600 for row in matching['combinations']))
        self.batch([quota_invoice], quota_id=quota['id'])
        exhausted = self.match(direction='below', quota_id=quota['id'])
        self.assertEqual(exhausted['combinations'], [])
        self.assertIn('剩余额度', exhausted['message'])
        self.assertEqual(self.match(1, direction='below')['combinations'], [])
        self.call('POST', '/api/match', {'target_cents': 1000, 'direction': 'invalid'}, status=400)

    def test_cancel_and_reject_release_invoices_but_keep_history_paid_is_final(self):
        invoice = self.invoice(seller='最初商家')
        first = self.batch([invoice])
        self.call('PATCH', '/api/batches/' + first['id'], {'status': 'cancelled'})
        self.call('PATCH', '/api/invoices/' + invoice['id'], {'seller': '修正商家'})
        self.call('DELETE', '/api/invoices/' + invoice['id'], status=409)
        historical = next(item for item in self.state()['batches'] if item['id'] == first['id'])
        self.assertEqual(historical['items'][0]['seller'], '最初商家')
        self.assertEqual(historical['status'], 'cancelled')
        second = self.batch([invoice])
        self.call('PATCH', '/api/batches/' + second['id'], {'status': 'submitted'})
        self.call('PATCH', '/api/batches/' + second['id'], {'status': 'rejected'})
        self.assertEqual(self.match()['candidate_count'], 1)
        third = self.batch([invoice])
        self.call('PATCH', '/api/batches/' + third['id'], {'status': 'submitted'})
        self.call('PATCH', '/api/batches/' + third['id'], {'status': 'paid'})
        self.assertEqual(self.match()['candidate_count'], 0)
        self.assertEqual(self.state()['summary']['paid_cents'], 1000)
        self.call('PATCH', '/api/batches/' + third['id'], {'status': 'cancelled'}, status=409)
        self.call('PATCH', '/api/invoices/' + invoice['id'], {'seller': '不能改'}, status=409)

    def test_monthly_quota_reserves_releases_and_rejects_overspend(self):
        quota = self.call('POST', '/api/quotas', {'month': '2026-10', 'category': '餐饮', 'amount_cents': 1000})
        first, second, other = self.invoice(600), self.invoice(500), self.invoice(200, category='交通')
        matching = self.match(quota_id=quota['id'])
        self.assertEqual(matching['candidate_count'], 2)
        self.assertEqual(matching['combinations'], [])
        self.assertIn('额度不足', matching['message'])
        self.call('POST', '/api/batches', {'name': '类别冲突', 'invoice_ids': [other['id']], 'quota_id': quota['id']}, status=400)
        batch = self.batch([first], quota_id=quota['id'])
        active = self.state()['quotas'][0]
        self.assertEqual((active['used_cents'], active['remaining_cents']), (600, 400))
        self.call('POST', '/api/batches', {'name': '超额', 'invoice_ids': [second['id']], 'quota_id': quota['id']}, status=409)
        self.call('POST', '/api/match', {'target_cents': 500, 'quota_id': quota['id']}, status=400)
        self.call('PATCH', '/api/batches/' + batch['id'], {'status': 'cancelled'})
        self.assertEqual(self.state()['quotas'][0]['used_cents'], 0)
        rejected = self.batch([second], quota_id=quota['id'])
        self.call('PATCH', '/api/batches/' + rejected['id'], {'status': 'submitted'})
        self.call('PATCH', '/api/batches/' + rejected['id'], {'status': 'rejected'})
        self.assertEqual(self.state()['quotas'][0]['remaining_cents'], 1000)
        paid = self.batch([first], quota_id=quota['id'])
        self.call('PATCH', '/api/batches/' + paid['id'], {'status': 'submitted'})
        self.call('PATCH', '/api/batches/' + paid['id'], {'status': 'paid'})
        self.assertEqual(self.state()['quotas'][0]['used_cents'], 600)

    def test_exports_include_original_and_literal_snapshot_excel_values(self):
        invoice, original = self.xml_import()
        invoice = self.call('PATCH', '/api/invoices/' + invoice['id'], {
            'verified': True, 'seller': '=SUM(1,2)', 'note': '=HYPERLINK("https://example.test","click")',
        })
        batch = self.batch([invoice])
        self.call('PATCH', '/api/batches/' + batch['id'], {'status': 'cancelled'})
        self.call('PATCH', '/api/invoices/' + invoice['id'], {'seller': '当前商家', 'amount_cents': 4321})
        response = self.client.get('/api/batches/' + batch['id'] + '/export')
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(archive.read('原件/原始发票.xml'), original)
            snapshot = json.loads(archive.read('报销单.json'))
            self.assertEqual(snapshot['items'][0]['amount_cents'], 1234)
            sheet = load_workbook(io.BytesIO(archive.read('发票明细.xlsx')), data_only=False).active
            self.assertEqual(sheet['A2'].value, '001234567890')
            self.assertEqual(sheet['B2'].value, '=SUM(1,2)')
            self.assertEqual(sheet['B2'].data_type, 's')
            self.assertEqual(sheet['K2'].data_type, 's')
            self.assertEqual(sheet['F2'].value, 12.34)
            self.assertEqual(sheet['H2'].value, '已取消')
        library = load_workbook(io.BytesIO(self.client.get('/api/export').content)).active
        self.assertEqual(library['B2'].value, '当前商家')
        self.assertEqual(library['F2'].value, 43.21)

    def test_original_download_and_editable_batch_filenames_use_snapshot(self):
        from urllib.parse import unquote
        invoice, original = self.xml_import()
        self.call('PATCH', '/api/invoices/' + invoice['id'], {'verified': True, 'category': '餐饮'})
        batch = self.batch([invoice])
        file_url = batch['items'][0]['file_url']
        response = self.client.get(invoice['file_url'], params={'download': 'true'})
        self.assertEqual(response.content, original)
        self.assertIn('attachment;', response.headers['content-disposition'])
        self.assertIn('原始发票.xml', unquote(response.headers['content-disposition']))
        self.assertTrue(self.client.get(file_url).headers['content-type'].startswith('application/xml'))
        self.call('PATCH', '/api/batches/' + batch['id'], {'status': 'cancelled'})
        self.call('PATCH', '/api/invoices/' + invoice['id'], {'category': '交通', 'amount_cents': 9999})
        export_url = '/api/batches/' + batch['id'] + '/export'
        response = self.client.post(export_url, json={'originals_only': True, 'rename': True})
        self.assertEqual(response.status_code, 200, response.text)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(archive.namelist(), ['原件/001-餐饮-12.34.xml'])
            self.assertEqual(archive.read(archive.namelist()[0]), original)
        options = {'originals_only': True, 'rename': True, 'filenames': {invoice['id']: '001-午餐-12.34'}}
        response = self.client.post(export_url, json=options)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(archive.namelist(), ['原件/001-午餐-12.34.xml'])
        response = self.client.get(file_url, params={'download': True, 'filename': '001-午餐-12.34'})
        self.assertEqual(response.content, original)
        self.assertIn('001-午餐-12.34.xml', unquote(response.headers['content-disposition']))
        options['rename'] = False
        response = self.client.post(export_url, json=options)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(archive.namelist(), ['原件/原始发票.xml'])
        self.assertEqual(self.state()['invoices'][0]['file_name'], '原始发票.xml')

    def test_batch_download_handles_duplicate_names_missing_files_and_invalid_names(self):
        first, original = self.xml_import('FIRST')
        second, other = self.xml_import('SECOND')
        for invoice in (first, second):
            self.call('PATCH', '/api/invoices/' + invoice['id'], {'verified': True})
        manual = self.invoice()
        batch = self.batch([first, second, manual])
        url = '/api/batches/' + batch['id'] + '/export'
        response = self.client.post(url, json={'originals_only': True})
        self.assertEqual(response.status_code, 200, response.text)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(set(archive.namelist()), {'原件/原始发票.xml', '原件/原始发票-2.xml', '无原件记录.txt'})
            self.assertEqual(archive.read('原件/原始发票.xml'), original)
            self.assertEqual(archive.read('原件/原始发票-2.xml'), other)
            self.assertIn('003_', archive.read('无原件记录.txt').decode())
        for bad in ('../escape', 'bad\\name', '', 'x' * 181):
            response = self.client.post(url, json={'rename': True, 'filenames': {first['id']: bad}})
            self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(self.client.post(url, json={'rename': True, 'filenames': {'outside': '文件'}}).status_code, 400)
        self.assertEqual(self.client.get(batch['items'][0]['file_url'], params={'download': True, 'filename': '../escape'}).status_code, 400)
        self.assertEqual(self.client.get(batch['items'][2]['file_url'] or '/api/batches/' + batch['id'] + '/invoices/' + manual['id'] + '/file').status_code, 404)

    def test_pdf_preview_renders_one_page_at_a_time_and_validates_page_range(self):
        import pymupdf
        invoice = self.invoice()
        relative = 'originals/' + invoice['id'] + '.pdf'
        with pymupdf.open() as document:
            for text in ('SAMPLE PAGE ONE', 'SAMPLE PAGE TWO'):
                document.new_page().insert_text((40, 40), text)
            document.save(self.data / relative)
        with db.connection(write=True) as conn:
            conn.execute('UPDATE invoices SET file_path=?,file_name=? WHERE id=?', (relative, 'sample.pdf', invoice['id']))
        batch = self.batch([invoice])
        for url in ('/api/invoices/' + invoice['id'] + '/file', batch['items'][0]['file_url']):
            for page in (1, 2):
                response = self.client.get(url, params={'preview_page': page})
                self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
                self.assertEqual(response.headers['content-type'], 'image/png')
                self.assertEqual(response.headers['x-preview-pages'], '2')
                image = pymupdf.Pixmap(response.content)
                self.assertLessEqual(max(image.width, image.height), 2000)
            for page in (0, -1, 3):
                self.assertEqual(self.client.get(url, params={'preview_page': page}).status_code, 400)
        xml, _ = self.xml_import()
        self.assertEqual(self.client.get(xml['file_url'], params={'preview_page': 1}).status_code, 400)

    def test_personal_backup_contains_own_records_and_originals_only(self):
        invoice, original = self.xml_import()
        self.call('PATCH', '/api/invoices/' + invoice['id'], {'verified': True})
        batch = self.batch([invoice])
        response = self.client.get('/api/backup')
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            payload = json.loads(archive.read('个人数据.json'))
            self.assertEqual(payload['format'], 'invoice-desk-personal-v2')
            self.assertEqual(payload['invoices'][0]['batch_id'], batch['id'])
            self.assertEqual(payload['batches'][0]['invoice_ids'], [invoice['id']])
            self.assertEqual(archive.read('原件/' + invoice['id'] + '_原始发票.xml'), original)
            self.assertFalse(any(name.endswith('.sqlite3') or 'secret.key' in name for name in archive.namelist()))

    def test_frontend_assets_revalidate_after_updates(self):
        response = self.client.get('/')
        self.assertEqual(response.headers.get('Cache-Control'), 'no-cache')
        self.assertEqual(self.client.get('/index.html').headers.get('Cache-Control'), 'no-cache')
        assets = re.findall(r'(?:href|src)="(/(?:style\.css|app\.js)\?v=[^"]+)"', response.text)
        self.assertEqual(len(assets), 2)
        for asset in assets:
            with self.subTest(asset=asset):
                response = self.client.get(asset)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.headers.get('Cache-Control'), 'no-cache')
                cached = self.client.get(asset, headers={'If-None-Match': response.headers['etag']})
                self.assertEqual(cached.status_code, 304)
                self.assertEqual(cached.headers.get('Cache-Control'), 'no-cache')

    def test_mutations_require_local_application_headers(self):
        self.assertEqual(self.client.post('/api/projects', json={'name': '拦截'}, headers={'X-Requested-With': ''}).status_code, 403)
        self.assertEqual(self.client.post('/api/projects', json={'name': '拦截'}, headers={'Origin': 'https://other.example'}).status_code, 403)
        self.assertEqual(self.client.get('/api/state', headers={'Host': 'other.example'}).status_code, 400)
        self.assertEqual(self.state()['projects'], [])
        self.assertEqual(self.client.get('/api/state').headers['Cache-Control'], 'no-store')


if __name__ == '__main__':
    unittest.main()
