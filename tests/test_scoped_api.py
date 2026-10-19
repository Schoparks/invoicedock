"""Multi-account boundaries and a single reservation across overlapping invoice pools."""
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from fastapi.testclient import TestClient
from app import db, main


class ScopedApiTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        data = Path(temporary.name)
        for module in (db, main):
            p = patch.object(module, 'DATA', data)
            p.start()
            self.addCleanup(p.stop)
        p = patch.dict(os.environ, {'INVOICE_MAIL_WORKER':'0', 'ALLOW_REGISTRATION':'true', 'COOKIE_SECURE':'false'})
        p.start()
        self.addCleanup(p.stop)
        self.a = TestClient(main.app, headers={'X-Requested-With':'InvoiceDesk'})
        self.a.__enter__()
        self.addCleanup(self.a.__exit__, None, None, None)
        self.b = TestClient(main.app, headers={'X-Requested-With':'InvoiceDesk'})
        self.c = TestClient(main.app, headers={'X-Requested-With':'InvoiceDesk'})
        self.addCleanup(self.b.close)
        self.addCleanup(self.c.close)
        for n, client in enumerate((self.a, self.b, self.c)):
            self.ok(client.post('/api/auth/register', json={'email':f'user{n}@example.test','name':f'用户{n}','password':'test-password-123'}))
        self.group = self.ok(self.a.post('/api/groups', json={'name':'出差组'}))
        self.group2 = self.ok(self.a.post('/api/groups', json={'name':'项目组'}))
        token = self.ok(self.a.post(f"/api/groups/{self.group['id']}/invites"))['token']
        self.ok(self.b.post('/api/groups/join', json={'token':token}))

    def ok(self, response):
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def upload(self):
        raw = b'<Invoice><InvoiceNumber>SHARED001</InvoiceNumber><SellerName>Vendor</SellerName><InvoiceDate>2026-10-15</InvoiceDate><TotalTaxIncludedAmount>12.34</TotalTaxIncludedAmount></Invoice>'
        item = self.ok(self.a.post('/api/import', files={'files':('private.xml',raw,'application/xml')}))['imported'][0]
        self.ok(self.a.patch('/api/invoices/'+item['id'],json={'verified':True}))
        return item

    def test_selective_sharing_access_and_global_reservation(self):
        invoice = self.upload()
        gid, other_gid = self.group['id'], self.group2['id']
        self.assertEqual(self.b.get(invoice['file_url']).status_code,403)
        self.assertEqual(self.ok(self.b.get('/api/state'))['invoices'],[])
        self.ok(self.a.put('/api/invoices/'+invoice['id']+'/shares',json={'group_ids':[gid,other_gid]}))
        shared = self.ok(self.b.get('/api/state',params={'group_id':gid}))['invoices'][0]
        self.assertFalse(shared['can_edit'])
        self.assertEqual(shared['shared_group_ids'],[])
        self.assertEqual(self.b.get(invoice['file_url']).status_code,200)
        self.assertEqual(self.b.patch('/api/invoices/'+invoice['id'],json={'seller':'forbidden'}).status_code,403)
        batch = self.ok(self.b.post('/api/batches',json={'name':'组报销','group_id':gid,'invoice_ids':[invoice['id']]}))
        self.assertEqual(self.ok(self.a.post('/api/match',json={'target_cents':2000,'group_id':other_gid}))['candidate_count'],0)
        self.assertEqual(self.a.post('/api/batches',json={'name':'重复使用','invoice_ids':[invoice['id']]}).status_code,409)
        self.assertEqual(self.a.put('/api/invoices/'+invoice['id']+'/shares',json={'group_ids':[other_gid]}).status_code,409)
        self.assertEqual(self.c.get('/api/state',params={'group_id':gid}).status_code,403)
        self.assertEqual(self.c.get('/api/batches/'+batch['id']+'/export').status_code,403)
        batch_file = batch['items'][0]['file_url']
        self.assertEqual(self.c.get(batch_file, params={'download': True}).status_code,403)
        self.assertEqual(self.c.get(batch_file, params={'preview_page': 1}).status_code,403)
        self.assertEqual(self.c.post('/api/batches/'+batch['id']+'/export', json={'originals_only': True}).status_code,403)
        other = self.ok(self.a.post('/api/invoices', json={'seller':'Other','date':'2026-10-15','amount_cents':100,'verified':True}))
        self.assertEqual(self.b.get('/api/batches/'+batch['id']+'/invoices/'+other['id']+'/file').status_code,404)
        self.ok(self.b.patch('/api/batches/'+batch['id'],json={'status':'cancelled'}))
        self.ok(self.a.put('/api/invoices/'+invoice['id']+'/shares',json={'group_ids':[other_gid]}))
        self.assertEqual(self.ok(self.b.get('/api/state',params={'group_id':gid}))['invoices'],[])
        self.assertEqual(self.b.get(invoice['file_url']).status_code,403)
        self.assertEqual(self.b.get(batch_file).status_code,200)
        exported = self.b.post('/api/batches/'+batch['id']+'/export', json={'originals_only': True, 'rename': True})
        self.assertEqual(exported.status_code,200)
        with zipfile.ZipFile(io.BytesIO(exported.content)) as archive:
            self.assertEqual(archive.namelist(), ['原件/001-其他-12.34.xml'])
        self.assertEqual(self.ok(self.a.post('/api/match',json={'target_cents':1234,'group_id':other_gid}))['combinations'][0]['total_cents'],1234)

    def test_leave_cancels_mixed_pending_batch_and_preserves_other_team_shares(self):
        gid = self.group['id']
        original = self.upload()
        member = self.ok(self.b.get('/api/state'))['user']
        own = self.ok(self.b.post('/api/invoices', json={'seller':'成员商家','date':'2026-10-15','amount_cents':500,'verified':True}))
        self.ok(self.a.put('/api/invoices/'+original['id']+'/shares', json={'group_ids':[gid,self.group2['id']]}))
        self.ok(self.b.put('/api/invoices/'+own['id']+'/shares', json={'group_ids':[gid]}))
        quota = self.ok(self.a.post('/api/quotas', json={'month':'2026-10','category':'全部','amount_cents':10000,'group_id':gid}))
        batch = self.ok(self.a.post('/api/batches', json={'name':'混合报销','group_id':gid,'quota_id':quota['id'],'invoice_ids':[original['id'],own['id']]}))
        self.ok(self.a.patch('/api/batches/'+batch['id'], json={'status':'submitted'}))
        response = self.ok(self.b.delete(f"/api/groups/{gid}/members/{member['id']}"))
        self.assertEqual(response['cancelled_batches'], 1)
        self.assertEqual(self.b.get('/api/state', params={'group_id':gid}).status_code, 403)
        self.assertEqual(self.b.get(original['file_url']).status_code, 403)
        returned = self.ok(self.b.get('/api/state'))['invoices'][0]
        self.assertEqual((returned['id'],returned['owner_id'],returned['status'],returned['shared_group_ids']), (own['id'],member['id'],'available',[]))
        team = self.ok(self.a.get('/api/state', params={'group_id':gid}))
        self.assertEqual(team['batches'][0]['status'], 'cancelled')
        self.assertEqual(team['batches'][0]['items'], batch['items'])
        self.assertEqual(team['quotas'][0]['remaining_cents'], 10000)
        self.assertEqual([i['id'] for i in team['invoices']], [original['id']])
        other = self.ok(self.a.get('/api/state', params={'group_id':self.group2['id']}))['invoices'][0]
        self.assertEqual((other['id'],other['status']), (original['id'],'available'))
        self.assertEqual(self.a.get(original['file_url']).status_code, 200)

    def test_kick_cancels_departing_creators_batches_but_preserves_paid_claims(self):
        gid = self.group['id']
        member = self.ok(self.b.get('/api/state'))['user']
        original = self.upload()
        paid = self.ok(self.b.post('/api/invoices', json={'seller':'已报销商家','date':'2026-10-15','amount_cents':600,'verified':True}))
        self.ok(self.a.put('/api/invoices/'+original['id']+'/shares', json={'group_ids':[gid]}))
        self.ok(self.b.put('/api/invoices/'+paid['id']+'/shares', json={'group_ids':[gid]}))
        draft = self.ok(self.b.post('/api/batches', json={'name':'离队成员创建','group_id':gid,'invoice_ids':[original['id']]}))
        settled = self.ok(self.a.post('/api/batches', json={'name':'已完成报销','group_id':gid,'invoice_ids':[paid['id']]}))
        self.ok(self.a.patch('/api/batches/'+settled['id'], json={'status':'submitted'}))
        self.ok(self.a.patch('/api/batches/'+settled['id'], json={'status':'paid'}))
        self.ok(self.a.delete(f"/api/groups/{gid}/members/{member['id']}"))
        team = self.ok(self.a.get('/api/state', params={'group_id':gid}))
        self.assertEqual({b['id']:b['status'] for b in team['batches']}, {draft['id']:'cancelled',settled['id']:'paid'})
        own = self.ok(self.b.get('/api/state'))['invoices'][0]
        self.assertEqual((own['id'],own['status'],own['batch_id'],own['shared_group_ids']), (paid['id'],'paid',settled['id'],[]))
        self.assertEqual(self.b.post('/api/batches', json={'name':'重复报销','invoice_ids':[paid['id']]}).status_code, 409)
        self.assertEqual(self.ok(self.a.get('/api/state'))['invoices'][0]['status'], 'available')

    def test_dissolve_preserves_originals_private_history_and_other_team(self):
        gid = self.group['id']
        original = self.upload()
        member = self.ok(self.b.get('/api/state'))['user']
        own = self.ok(self.b.post('/api/invoices', json={'seller':'成员商家','date':'2026-10-15','amount_cents':500,'verified':True}))
        self.ok(self.a.put('/api/invoices/'+original['id']+'/shares', json={'group_ids':[gid,self.group2['id']]}))
        self.ok(self.b.put('/api/invoices/'+own['id']+'/shares', json={'group_ids':[gid]}))
        project = self.ok(self.a.post('/api/projects', json={'name':'团队项目','group_id':gid}))
        quota = self.ok(self.a.post('/api/quotas', json={'month':'2026-10','category':'全部','amount_cents':10000,'group_id':gid}))
        self.ok(self.a.patch('/api/invoices/'+original['id'], json={'project_id':project['id']}))
        history = self.ok(self.b.post('/api/batches', json={'name':'历史报销','group_id':gid,'quota_id':quota['id'],'project_id':project['id'],'invoice_ids':[original['id']]}))
        self.ok(self.b.patch('/api/batches/'+history['id'], json={'status':'submitted'}))
        self.ok(self.b.patch('/api/batches/'+history['id'], json={'status':'paid'}))
        draft = self.ok(self.a.post('/api/batches', json={'name':'未完成报销','group_id':gid,'invoice_ids':[own['id']]}))
        invite = self.ok(self.a.post(f'/api/groups/{gid}/invites'))['token']
        self.assertEqual(self.b.delete(f'/api/groups/{gid}').status_code, 403)
        self.assertEqual(self.c.delete(f'/api/groups/{gid}').status_code, 403)
        response = self.ok(self.a.delete(f'/api/groups/{gid}'))
        self.assertEqual(response['cancelled_batches'], 1)
        self.assertEqual(self.c.post('/api/groups/join', json={'token':invite}).status_code, 400)
        self.assertEqual(self.a.get('/api/state', params={'group_id':gid}).status_code, 403)
        personal_a, personal_b = self.ok(self.a.get('/api/state')), self.ok(self.b.get('/api/state'))
        self.assertEqual([g['id'] for g in personal_a['groups']], [self.group2['id']])
        self.assertEqual(personal_b['groups'], [])
        self.assertEqual((personal_a['invoices'][0]['id'],personal_a['invoices'][0]['status']), (original['id'],'paid'))
        self.assertEqual((personal_b['invoices'][0]['id'],personal_b['invoices'][0]['owner_id'],personal_b['invoices'][0]['status']), (own['id'],member['id'],'available'))
        self.assertIsNone(personal_a['invoices'][0]['project_id'])
        self.assertEqual(personal_b['batches'][0]['items'], history['items'])
        self.assertEqual(personal_b['batches'][0]['status'], 'paid')
        self.assertEqual(personal_a['batches'][0]['status'], 'cancelled')
        self.assertEqual(self.b.get(history['items'][0]['file_url']).status_code, 200)
        self.assertEqual(self.c.get(history['items'][0]['file_url']).status_code, 403)
        self.assertEqual(self.a.get('/api/batches/'+history['id']+'/export').status_code, 403)
        self.assertEqual(self.a.get(original['file_url']).status_code, 200)
        self.assertEqual(self.b.get(original['file_url']).status_code, 403)
        self.assertEqual(self.a.get('/api/batches/'+draft['id']+'/export').status_code, 200)
        with db.connection() as connection:
            self.assertEqual(connection.execute('PRAGMA foreign_key_check').fetchall(), [])
            self.assertIsNone(connection.execute('SELECT id FROM projects WHERE id=?', (project['id'],)).fetchone())
            self.assertIsNone(connection.execute('SELECT id FROM quotas WHERE id=?', (quota['id'],)).fetchone())
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM batch_items').fetchone()[0], 2)

    def test_personal_export_and_scoped_quotas_do_not_leak(self):
        invoice = self.upload()
        gid = self.group['id']
        quota = {'month':'2026-10','category':'餐饮','amount_cents':10000}
        self.ok(self.a.post('/api/quotas',json=quota))
        self.ok(self.b.post('/api/quotas',json=quota))
        self.ok(self.a.post('/api/quotas',json={**quota,'group_id':gid}))
        self.assertEqual(self.b.post('/api/quotas',json={**quota,'category':'交通','group_id':gid}).status_code,403)
        self.ok(self.a.put('/api/invoices/'+invoice['id']+'/shares',json={'group_ids':[gid]}))
        backup=self.b.get('/api/backup')
        self.assertEqual(backup.status_code,200)
        with zipfile.ZipFile(io.BytesIO(backup.content)) as archive:
            content=json.loads(archive.read('个人数据.json'))
            self.assertEqual(content['invoices'],[])
            self.assertFalse(any(n.startswith('原件/') for n in archive.namelist()))
        self.ok(self.b.post('/api/auth/logout'))
        self.assertEqual(self.b.get('/api/state').status_code,401)
        self.assertEqual(self.b.get(invoice['file_url']).status_code,401)

    def test_unified_view_keeps_account_and_share_boundaries(self):
        invoice = self.upload()
        gid = self.group['id']
        self.ok(self.a.put('/api/invoices/'+invoice['id']+'/shares', json={'group_ids':[gid,self.group2['id']]}))
        project = self.ok(self.a.post('/api/projects', json={'name':'组项目','group_id':gid}))
        unified = self.ok(self.b.get('/api/state', params={'unified':'true'}))
        self.assertEqual([i['id'] for i in unified['invoices']], [invoice['id']])
        self.assertEqual(unified['invoices'][0]['groups'], [{'id':gid,'name':'出差组'}])
        self.assertEqual([p['id'] for p in unified['projects']], [project['id']])
        self.assertEqual(len(self.ok(self.a.get('/api/state', params={'unified':'true'}))['invoices']), 1)
        self.assertEqual(self.ok(self.c.get('/api/state', params={'unified':'true'}))['invoices'], [])
        self.assertEqual(self.ok(self.b.get('/api/state'))['invoices'], [])
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(self.b.get('/api/export?unified=true').content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual((sheet['L2'].value, sheet['M2'].value), ('出差组', '用户0'))
        self.assertEqual(self.c.post('/api/batches', json={'name':'越权','unified':True,'invoice_ids':[invoice['id']]}).status_code, 403)
        self.ok(self.a.put('/api/invoices/'+invoice['id']+'/shares', json={'group_ids':[]}))
        self.assertEqual(self.ok(self.b.get('/api/state', params={'unified':'true'}))['invoices'], [])
        self.assertEqual(self.b.post('/api/batches', json={'name':'失效分享','unified':True,'invoice_ids':[invoice['id']]}).status_code, 403)

    def test_unified_match_priority_and_mixed_batch(self):
        def add(client, amount, group=None):
            return self.ok(client.post('/api/invoices', json={'seller':'合成商家','date':'2026-10-19','amount_cents':amount,'verified':True,'group_id':group}))
        gid = self.group['id']
        personal = add(self.b, 1200)
        owned = add(self.b, 1200, gid)
        shared = add(self.a, 1200, gid)
        exact = add(self.a, 1100, gid)
        def match(target, **extra):
            return self.ok(self.b.post('/api/match', json={'target_cents':target,'unified':True,**extra}))['combinations'][0]
        self.assertEqual(match(1100)['invoice_ids'], [exact['id']])
        self.assertEqual(match(1200)['invoice_ids'], [personal['id']])
        self.assertEqual(match(1200, group_id=gid)['invoice_ids'], [owned['id']])
        self.assertEqual(match(1200, source='personal')['invoice_ids'], [personal['id']])
        batch = self.ok(self.b.post('/api/batches', json={'name':'统一报销','unified':True,'invoice_ids':[personal['id'],shared['id']]}))
        self.assertEqual(batch['total_cents'], 2400)
        self.assertIsNone(batch['group_id'])
        self.assertEqual(self.a.get('/api/batches/'+batch['id']+'/export').status_code, 403)
        self.assertEqual(match(1200, group_id=gid)['invoice_ids'], [owned['id']])
        quota = self.ok(self.a.post('/api/quotas', json={'month':'2026-10','category':'全部','amount_cents':5000,'group_id':gid}))
        result = self.ok(self.b.post('/api/match', json={'unified':True,'target_cents':1200,'quota_id':quota['id']}))
        self.assertEqual(result['group_id'], gid)
        self.assertEqual(result['combinations'][0]['invoice_ids'], [owned['id']])
        personal_batch = self.b.post('/api/batches', json={'name':'错误额度归属','unified':True,'invoice_ids':[owned['id']],'quota_id':quota['id']})
        self.assertEqual(personal_batch.status_code, 400)
        personal_quota = self.ok(self.b.post('/api/quotas', json={'month':'2026-10','category':'全部','amount_cents':5000}))
        result = self.ok(self.b.post('/api/match', json={'unified':True,'target_cents':1200,'group_id':gid,'quota_id':personal_quota['id']}))
        self.assertIsNone(result['group_id'])
        self.assertEqual(result['combinations'][0]['invoice_ids'], [owned['id']])


if __name__ == '__main__':
    unittest.main()
